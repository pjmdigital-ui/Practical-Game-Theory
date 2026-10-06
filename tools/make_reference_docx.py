#!/usr/bin/env python3
"""Create tools/reference.docx: the Word style template used by build_book.py
(6x9 trim, Georgia, page-number footer, each chapter on a new page)."""
import subprocess
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "tools/reference.docx"
with open(OUT, "wb") as f:
    f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                           capture_output=True, check=True).stdout)
d = Document(OUT)
for s in d.sections:
    s.page_width, s.page_height = Inches(6), Inches(9)
    s.left_margin = s.right_margin = Inches(0.75)
    s.top_margin = s.bottom_margin = Inches(0.85)
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    r.font.size = Pt(9)

styles = {s.name: s for s in d.styles}


def setf(name, size, bold=None, color=(0x1B, 0x1B, 0x1B), font="Georgia"):
    s = styles.get(name)
    if s is None:
        return
    s.font.name, s.font.size = font, Pt(size)
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    if bold is not None:
        s.font.bold = bold
    s.font.color.rgb = RGBColor(*color)


for n in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    setf(n, 11)
setf("Title", 28, True); setf("Subtitle", 14, False, (0x44, 0x44, 0x44)); setf("Author", 12)
setf("Heading 1", 22, True); setf("Heading 2", 14, True); setf("Heading 3", 11.5, True)
setf("TOC Heading", 18, True)
h1 = styles["Heading 1"].paragraph_format
h1.page_break_before, h1.space_before, h1.space_after = True, Pt(72), Pt(30)
for n in ("Body Text", "First Paragraph"):
    pf = styles[n].paragraph_format
    pf.space_after, pf.line_spacing = Pt(6), 1.2
d.save(OUT)
print("wrote", OUT)
